from io import BytesIO
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from fpdf import FPDF

from utils.text_utils import sanitize_text


def split_sections(text):

    lines = sanitize_text(text).splitlines()

    sections = []

    current_title = None
    current_lines = []

    for line in lines:

        stripped = line.strip()

        is_heading = (
            stripped
            and len(stripped) < 100
            and (
                stripped.isupper()
                or stripped.endswith(":")
                or re.match(
                    r"^(PARTIES|PURPOSE|TERMS|GENERAL|SIGNATURES|REVIEW NOTICE)",
                    stripped,
                    re.I
                )
            )
        )

        if is_heading:

            if current_title is not None:

                sections.append(
                    (
                        current_title,
                        current_lines
                    )
                )

            current_title = stripped.rstrip(":")
            current_lines = []

        else:

            current_lines.append(
                stripped
            )

    if current_title is not None:

        sections.append(
            (
                current_title,
                current_lines
            )
        )

    else:

        sections = [
            ("", lines)
        ]

    return sections


def format_txt(text):

    return sanitize_text(
        text
    ).encode("utf-8")


def format_docx(
    text,
    document_type
):

    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal_style = document.styles["Normal"]

    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(11)

    # Main title

    paragraph = document.add_paragraph()

    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = paragraph.add_run(
        "LegalEase"
    )

    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(22)

    # Document type

    paragraph = document.add_paragraph()

    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = paragraph.add_run(
        document_type
    )

    run.bold = True
    run.font.size = Pt(14)

    # Document content

    sections = split_sections(text)

    for heading, lines in sections:

        if heading:

            paragraph = document.add_paragraph()

            run = paragraph.add_run(
                heading
            )

            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)

        for line in lines:

            if not line:
                continue

            paragraph = document.add_paragraph()

            if line.startswith("- "):

                paragraph.style = (
                    document.styles["List Bullet"]
                )

                line = line[2:]

            paragraph.add_run(
                line
            )

    # Footer

    footer = section.footer.paragraphs[0]

    footer.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    footer_run = footer.add_run(
        "LegalEase - Generated legal-information draft"
    )

    footer_run.italic = True

    output = BytesIO()

    document.save(output)

    return output.getvalue()


class LegalEasePDF(FPDF):

    def __init__(self, document_type):

        super().__init__()

        self.document_type = document_type

        self.set_auto_page_break(
            auto=True,
            margin=18
        )

    def header(self):

        self.set_font(
            "Helvetica",
            "B",
            14
        )

        self.cell(
            0,
            8,
            "LegalEase",
            ln=1,
            align="C"
        )

        self.set_font(
            "Helvetica",
            "B",
            11
        )

        self.cell(
            0,
            7,
            self.document_type,
            ln=1,
            align="C"
        )

        self.ln(3)

    def footer(self):

        self.set_y(-15)

        self.set_font(
            "Helvetica",
            "I",
            8
        )

        self.cell(
            0,
            8,
            "LegalEase - Generated legal-information draft",
            align="C"
        )


def format_pdf(
    text,
    document_type
):

    pdf = LegalEasePDF(
        document_type
    )

    pdf.add_page()

    pdf.set_font(
        "Helvetica",
        size=10.5
    )

    lines = sanitize_text(
        text
    ).splitlines()

    for raw_line in lines:

        line = raw_line.strip()

        if not line:

            pdf.ln(3)

            continue

        is_heading = (
            len(line) < 100
            and (
                line.isupper()
                or line.endswith(":")
            )
        )

        if is_heading:

            pdf.set_font(
                "Helvetica",
                "B",
                11
            )

            pdf.multi_cell(
                0,
                6,
                line.rstrip(":")
            )

            pdf.set_font(
                "Helvetica",
                size=10.5
            )

        elif line.startswith("- "):

            pdf.multi_cell(
                0,
                5.5,
                "- " + line[2:]
            )

        else:

            pdf.multi_cell(
                0,
                5.5,
                line
            )

    return bytes(
        pdf.output()
    )